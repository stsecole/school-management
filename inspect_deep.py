#!/usr/bin/env python3
"""Deep inspection: find last row with real data, dump all non-empty rows."""
import openpyxl
import warnings
warnings.filterwarnings("ignore")

FILES = [
    "/home/z/my-project/upload/متابعة المهام.xlsx",
    "/home/z/my-project/upload/قوائم-الحضور (4).xlsx",
    "/home/z/my-project/upload/تسجيلات-تقني-سامي-فيفري-2023.xlsx",
]


def last_nonempty_row(ws, max_scan=2_000_000):
    """Scan from bottom up to find the last row containing any non-empty cell."""
    max_r = min(ws.max_row, max_scan)
    for r in range(max_r, 0, -1):
        row = next(ws.iter_rows(min_row=r, max_row=r, values_only=True))
        if any(v not in (None, "") for v in row):
            return r
    return 0


def dump(ws, label):
    print(f"\n========== {label} ==========")
    last = last_nonempty_row(ws)
    print("last non-empty row:", last, " | max_row reported:", ws.max_row,
          " | max_col:", ws.max_column)
    print("\n-- All non-empty rows (row_idx: [values]) --")
    shown = 0
    for ridx, row in enumerate(ws.iter_rows(min_row=1, max_row=max(last, 1), values_only=True), start=1):
        if any(v not in (None, "") for v in row):
            cells = [("" if v is None else str(v)) for v in row]
            # Trim trailing empties for readability
            while cells and cells[-1] == "":
                cells.pop()
            print(f"  r{ridx}: {cells}")
            shown += 1
    print("total non-empty rows printed:", shown)


def main():
    for f in FILES:
        print("\n" + "#" * 110)
        print("# FILE:", f)
        print("#" * 110)
        wb = openpyxl.load_workbook(f, data_only=True)
        for ws in wb.worksheets:
            dump(ws, f"{f} :: {ws.title}")
        # Also try the values without data_only to see formulas
        wb2 = openpyxl.load_workbook(f, data_only=False)
        for ws in wb2.worksheets:
            last = last_nonempty_row(ws)
            if last == 0:
                continue
            print(f"\n  [formula view] sheet={ws.title!r} last_row={last}")
            for ridx, row in enumerate(ws.iter_rows(min_row=1, max_row=last, values_only=True), start=1):
                if any(v not in (None, "") for v in row):
                    has_formula = any(isinstance(v, str) and v.startswith("=") for v in row)
                    if has_formula:
                        cells = [("" if v is None else str(v)) for v in row]
                        while cells and cells[-1] == "":
                            cells.pop()
                        print(f"    r{ridx} (formula): {cells}")


if __name__ == "__main__":
    main()
